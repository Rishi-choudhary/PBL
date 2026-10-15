"""Builds the PBL report from the university template and content.py.

Keeps the template's cover, certificate, abstract, list-of-tables and list-of-figures pages (filled in),
removes the guideline pages, and adds the table of contents and chapters in the format the guidelines set:
A4, 0.7" margins, Times New Roman 12 justified at 1.15 spacing, headings bold 14, numbered italic
sub-headings, table captions above and figure captions below with chapter-wise numbering, page numbers
bottom-right starting at the Introduction, IEEE references in order of citation.

Two passes: the first build is converted to PDF to find the page of every heading and caption, the second
writes those page numbers into the table of contents and the lists of tables and figures.
Run from this folder: python3 build_report.py
"""
import copy
import re
import subprocess
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Twips
from docx.text.paragraph import Paragraph

import content as C

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "DSE_3170_PBL_3_Report_Format_Final.docx"
OUT = HERE / "Crop_Recommendation_System_PBL_Report.docx"
SOFFICE = Path("/root/.claude/skills/synced/05da9f5f-ce6b-4689-82b1-eee8c0b85ee9_d3fe8dfa-7d4e-4d1a-91dd-b8d3feeb5e35/"
               "docx/scripts/office/soffice.py")
MARGIN = 1008                      # 0.7 inch in twips
TEXT_W = 11900 - 2 * MARGIN        # template page width 11900 twips
FONT = "Times New Roman"
BULLET_NUM_ID = "14"               # template list: Symbol bullet, indent 720 / hanging 360
DECIMAL_ABSTRACT = "12"            # template list: "%1." decimal


# ----------------------------------------------------------------------------- numbering pre-pass
def number_items():
    """Assign section, figure, table and equation numbers in document order."""
    sec, fig, tab = {}, {}, {}
    headings, figures, tables = [], [], []
    ch = h2 = h3 = 0
    eq_count = {}
    eqs = []
    for item in C.CONTENT:
        kind = item[0]
        if kind == "h1":
            ch += 1; h2 = h3 = 0
            label = f"{ch}. {item[1]}"
            headings.append((1, label))
            if item[2]: sec[item[2]] = str(ch)
        elif kind == "h2":
            h2 += 1; h3 = 0
            num = f"{ch}.{h2}"
            headings.append((2, f"{num} {item[1]}"))
            if item[2]: sec[item[2]] = num
        elif kind == "h3":
            h3 += 1
            num = f"{ch}.{h2}.{h3}"
            headings.append((3, f"{num} {item[1]}"))
            if item[2]: sec[item[2]] = num
        elif kind == "appx":
            headings.append((2, f"{item[1]}. {item[2]}"))
        elif kind == "fig":
            n = sum(1 for f in figures if f[0].startswith(f"{ch}.")) + 1
            fig[item[1]] = f"{ch}.{n}"
            figures.append((f"{ch}.{n}", item[3]))
        elif kind == "tab":
            n = sum(1 for t in tables if t[0].startswith(f"{ch}.")) + 1
            tab[item[1]] = f"{ch}.{n}"
            tables.append((f"{ch}.{n}", item[2]))
        elif kind == "eq":
            eq_count[ch] = eq_count.get(ch, 0) + 1
            eqs.append(f"{ch}.{eq_count[ch]}")
    return sec, fig, tab, headings, figures, tables, eqs


SEC, FIGN, TABN, HEADINGS, FIGURES, TABLES, EQS = number_items()
CITE_ORDER = []


def resolve(text):
    def cite(m):
        out = []
        for k in m.group(1).split(","):
            k = k.strip()
            if k not in C.REFS:
                raise KeyError(f"unknown reference {k}")
            if k not in CITE_ORDER:
                CITE_ORDER.append(k)
            out.append(f"[{CITE_ORDER.index(k) + 1}]")
        return ", ".join(out)
    text = re.sub(r"\{c:([^}]+)\}", cite, text)
    text = re.sub(r"\{f:([^}]+)\}", lambda m: "Fig. " + FIGN[m.group(1)], text)
    text = re.sub(r"\{t:([^}]+)\}", lambda m: "Table " + TABN[m.group(1)], text)
    text = re.sub(r"\{s:([^}]+)\}", lambda m: SEC[m.group(1)], text)
    return text


# ----------------------------------------------------------------------------- run helpers
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\[\[.+?\]\]|_\{.+?\}|\^\{.+?\})")


def set_font(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def add_runs(par, text, size=12, bold=False, italic=False):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_runs(par, tok[2:-2], size, True, italic)
        elif tok.startswith("[[") and tok.endswith("]]"):
            r = par.add_run("[" + tok[2:-2] + "]")
            set_font(r, size, bold, italic)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif tok.startswith("_{") and tok.endswith("}"):
            r = par.add_run(tok[2:-1]); set_font(r, size, bold, italic); r.font.subscript = True
        elif tok.startswith("^{") and tok.endswith("}"):
            r = par.add_run(tok[2:-1]); set_font(r, size, bold, italic); r.font.superscript = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            add_runs(par, tok[1:-1], size, bold, True)
        else:
            r = par.add_run(tok); set_font(r, size, bold, italic)


def new_par(doc, style="Report Body", align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    return p


# ----------------------------------------------------------------------------- styles
def make_styles(doc):
    def para_style(name, size, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6,
                   line=1.15, outline=None, keep_next=False):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(size); st.font.bold = bold; st.font.italic = italic
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.alignment = align; pf.space_before = Pt(before); pf.space_after = Pt(after)
        pf.line_spacing = line; pf.keep_with_next = keep_next
        if outline is not None:
            lvl = OxmlElement("w:outlineLvl"); lvl.set(qn("w:val"), str(outline))
            st.element.get_or_add_pPr().append(lvl)
        return st
    para_style("Heading 1", 14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=12, outline=0, keep_next=True)
    para_style("Heading 2", 12, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, outline=1,
               keep_next=True)
    para_style("Heading 3", 12, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=6, outline=2,
               keep_next=True)
    para_style("Report Body", 12)
    para_style("Report List", 12, after=3)
    para_style("Report Caption", 11, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=10, line=1.0)
    para_style("Report Table", 11, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
    para_style("TOC Entry", 12, align=WD_ALIGN_PARAGRAPH.LEFT, after=4, line=1.15)
    para_style("Front Title", 16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12, line=1.5)


# ----------------------------------------------------------------------------- lists
def new_num(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    num_id = max(ids) + 1
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), abstract_id); num.append(a)
    ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
    numbering.append(num)
    return str(num_id)


def list_par(doc, text, num_id):
    p = new_par(doc, "Report List")
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); numpr.append(il)
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), num_id); numpr.append(ni)
    ppr.append(numpr)
    p.paragraph_format.left_indent = Twips(720)
    p.paragraph_format.first_line_indent = Twips(-360)
    add_runs(p, resolve(text))


# ----------------------------------------------------------------------------- tables
def set_cell_width(cell, w):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW"); tcpr.append(tcw)
    tcw.set(qn("w:w"), str(w)); tcw.set(qn("w:type"), "dxa")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcpr.append(sh)


def borders(table):
    """Table properties in schema order: tblStyle, tblW, jc, tblBorders, tblLayout, tblLook."""
    tblpr = table._tbl.tblPr
    style = tblpr.find(qn("w:tblStyle"))
    look = tblpr.find(qn("w:tblLook"))
    for child in list(tblpr):
        tblpr.remove(child)
    if style is not None:
        tblpr.append(style)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(TEXT_W)); w.set(qn("w:type"), "dxa"); tblpr.append(w)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblpr.append(jc)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "808080")
        b.append(e)
    tblpr.append(b)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    if look is not None:
        tblpr.append(look)


NUMERIC = re.compile(r"^[<−\-]?[\d.,]+( ± [\d.]+)?$")


def add_table(doc, key, caption, head, rows, weights, size):
    cap = new_par(doc, "Report Caption")
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(8); cap.paragraph_format.space_after = Pt(4)
    add_runs(cap, f"**Table {TABN[key]}:** {caption}", 11)
    total = sum(weights)
    widths = [int(TEXT_W * w / total) for w in weights]
    widths[-1] = TEXT_W - sum(widths[:-1])
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    borders(t)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for r_i, row in enumerate([head] + rows):
        tr = t.rows[r_i]
        trpr = tr._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trpr.append(cs)
        if r_i == 0:
            th = OxmlElement("w:tblHeader"); trpr.append(th)
        for c_i, txt in enumerate(row):
            cell = tr.cells[c_i]
            set_cell_width(cell, widths[c_i])
            p = cell.paragraphs[0]
            p.style = doc.styles["Report Table"]
            txt = resolve(txt)
            if r_i == 0:
                shade(cell, "D9D9D9")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, txt, size, bold=True)
            else:
                if c_i > 0 and NUMERIC.match(txt):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, txt, size)
    spacer = new_par(doc, "Report Body")
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 1.0
    for r in spacer.runs:
        r.font.size = Pt(4)


# ----------------------------------------------------------------------------- figures, equations
def add_figure(doc, key, path, caption, width):
    p = new_par(doc, "Report Body", WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(HERE / path), width=Inches(width))
    cap = new_par(doc, "Report Caption")
    add_runs(cap, f"**Figure {FIGN[key]}:** {caption}", 11)


EQ_ITER = iter(EQS)


def add_equation(doc, text):
    p = new_par(doc, "Report Body", WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Twips(TEXT_W // 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Twips(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    add_runs(p, "\t" + resolve(text) + "\t(" + next(EQ_ITER) + ")")


# ----------------------------------------------------------------------------- page numbers
def add_page_field(par):
    def fld(t):
        r = OxmlElement("w:r"); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r.append(e); return r
    rpr_run = par.add_run(""); set_font(rpr_run, 12)
    par._p.append(fld("begin"))
    r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    r.append(it); par._p.append(r)
    par._p.append(fld("separate"))
    num = par.add_run("1"); set_font(num, 12)
    par._p.append(fld("end"))


def set_margins(sectpr):
    pm = sectpr.find(qn("w:pgMar"))
    for side in ("top", "bottom", "left", "right"):
        pm.set(qn(f"w:{side}"), str(MARGIN))
    pm.set(qn("w:header"), "567"); pm.set(qn("w:footer"), "567")
    cols = sectpr.find(qn("w:cols"))
    if cols is not None:
        sectpr.remove(cols)
    cols = OxmlElement("w:cols"); cols.set(qn("w:space"), "720")
    sectpr.find(qn("w:pgMar")).addnext(cols)
    return cols


# ----------------------------------------------------------------------------- front matter
def body_children(doc):
    return list(doc.element.body.iterchildren())


def fill_front(doc, pages):
    body = body_children(doc)
    cover, cert, abstract_head = body[1], body[43], body[70]
    # cover: title and third student line (the text box is stored twice: Word and fallback copies)
    seen = 0
    for t in cover.iter(qn("w:t")):
        if t.text == "TITLE":
            t.text = C.TITLE_SHORT
        elif t.text == "_OF_PROJECT":
            t.text = ""
        elif t.text == "2 (Student2 ":
            if seen % 2 == 1:
                t.text = "3 (Student3 "
            seen += 1
    # certificate
    for t in cert.iter(qn("w:t")):
        if t.text in ("<", ">"):
            t.text = ""
        elif t.text == "PROJECT TITLE":
            t.text = C.TITLE_SHORT
    # abstract: replace the instructions with the two paragraphs
    for el in body[71:75]:
        el.getparent().remove(el)
    for text in reversed(C.ABSTRACT):
        p = doc.add_paragraph(style="Report Body")
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(10)
        add_runs(p, text)
        abstract_head.addnext(p._p)
    # lists of tables and figures
    lot_tbl, lof_head, lof_tbl = body[78], body[110], body[112]
    for el in body[79:110]:
        el.getparent().remove(el)
    ppr = lof_head.find(qn("w:pPr"))
    pb = OxmlElement("w:pageBreakBefore"); ppr.insert(0, pb)
    fill_list(doc, lot_tbl, [(n, cap, pages.get(f"Table {n}:", "")) for n, cap in TABLES])
    fill_list(doc, lof_tbl, [(n, cap, pages.get(f"Figure {n}:", "")) for n, cap in FIGURES])
    # guideline pages after the list of figures are removed
    for el in body[114:]:
        if el.tag != qn("w:sectPr"):
            el.getparent().remove(el)
    return body[113]


def fill_list(doc, tbl_el, entries):
    rows = tbl_el.findall(qn("w:tr"))
    proto = rows[1]
    for r in rows[1:]:
        tbl_el.remove(r)
    for num, cap, page in entries:
        tr = copy.deepcopy(proto)
        tbl_el.append(tr)
        for tc, txt in zip(tr.findall(qn("w:tc")), (num, cap, str(page))):
            for p in tc.findall(qn("w:p"))[1:]:
                tc.remove(p)
            p_el = tc.find(qn("w:p"))
            for r in p_el.findall(qn("w:r")):
                p_el.remove(r)
            par = Paragraph(p_el, None)
            if txt is cap:
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(par, txt, 11)


# ----------------------------------------------------------------------------- TOC
def add_toc(doc, pages, front_sectpr):
    p = new_par(doc, "Front Title")
    p.paragraph_format.page_break_before = False
    add_runs(p, "TABLE OF CONTENTS", 16, bold=True)
    for level, text in HEADINGS:
        e = new_par(doc, "TOC Entry")
        e.paragraph_format.left_indent = Twips({1: 0, 2: 400, 3: 800}[level])
        e.paragraph_format.space_before = Pt(6 if level == 1 else 0)
        e.paragraph_format.tab_stops.add_tab_stop(Twips(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        add_runs(e, text + "\t" + str(pages.get(text, "")), 12, bold=(level == 1))
    end = doc.add_paragraph()
    sp = copy.deepcopy(front_sectpr.find(".//" + qn("w:sectPr")))
    set_margins(sp)
    end._p.get_or_add_pPr().append(sp)


# ----------------------------------------------------------------------------- body
def add_body(doc):
    first_h1 = True
    after_appx = False
    num_cache = None
    for item in C.CONTENT:
        kind = item[0]
        if kind == "h1":
            label = next(t for lvl, t in HEADINGS if lvl == 1 and t.split(". ", 1)[1] == item[1])
            p = new_par(doc, "Heading 1")
            p.paragraph_format.page_break_before = not first_h1
            first_h1 = False
            add_runs(p, label, 14, bold=True)
        elif kind in ("h2", "h3"):
            num = SEC.get(item[2]) if item[2] else None
            p = new_par(doc, "Heading 2" if kind == "h2" else "Heading 3")
            add_runs(p, f"{num} {item[1]}", 12, italic=True)
        elif kind == "appx":
            p = new_par(doc, "Heading 2")
            add_runs(p, f"{item[1]}. {item[2]}", 12, italic=True)
            after_appx = True
        elif kind == "p":
            p = new_par(doc)
            add_runs(p, resolve(item[1]))
        elif kind == "bul":
            for it in item[1]:
                list_par(doc, it, BULLET_NUM_ID)
        elif kind == "num":
            nid = new_num(doc, DECIMAL_ABSTRACT)
            for it in item[1]:
                list_par(doc, it, nid)
        elif kind == "eq":
            add_equation(doc, item[1])
        elif kind == "tab":
            add_table(doc, *item[1:])
        elif kind == "fig":
            add_figure(doc, *item[1:])
        elif kind == "img":
            p = new_par(doc, "Report Body", WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.page_break_before = not after_appx
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(str(HERE / item[1]), width=Inches(item[2]))
            after_appx = False
        elif kind == "break":
            p = new_par(doc)
            p.paragraph_format.page_break_before = True
            after_appx = False
        elif kind == "refs":
            add_references(doc)
    return num_cache


def add_references(doc):
    for i, key in enumerate(CITE_ORDER, 1):
        p = new_par(doc, "Report Body", WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Twips(600)
        p.paragraph_format.first_line_indent = Twips(-600)
        p.paragraph_format.tab_stops.add_tab_stop(Twips(600))
        add_runs(p, f"[{i}]\t" + C.REFS[key])
    unused = [k for k in C.REFS if k not in CITE_ORDER]
    if unused:
        print("WARNING: references never cited:", unused)


# ----------------------------------------------------------------------------- build + pages
def build(pages, out):
    CITE_ORDER.clear()
    global EQ_ITER
    EQ_ITER = iter(EQS)
    doc = docx.Document(str(TEMPLATE))
    make_styles(doc)
    front_sectpr = fill_front(doc, pages)
    add_toc(doc, pages, front_sectpr)
    add_body(doc)
    last = doc.sections[-1]
    sp = last._sectPr
    cols = set_margins(sp)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
    cols.addprevious(pg)
    last.footer.is_linked_to_previous = False
    fp = last.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(fp)
    doc.core_properties.title = C.TITLE_SHORT.title()
    doc.save(str(out))


def find_pages(docx_path):
    pdf_dir = docx_path.parent / "_render"
    pdf_dir.mkdir(exist_ok=True)
    subprocess.run(["python3", str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir),
                    str(docx_path)], check=True, capture_output=True)
    pdf = pdf_dir / (docx_path.stem + ".pdf")
    text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True,
                          text=True).stdout
    page_lines = [[" ".join(l.split()) for l in pg.splitlines()] for pg in text.split("\f")]
    intro = HEADINGS[0][1]
    start = next(i for i, lines in enumerate(page_lines) if intro in lines)
    pages = {}
    for _, label in HEADINGS:
        for i in range(start, len(page_lines)):
            if label in page_lines[i]:
                pages[label] = i - start + 1
                break
        else:
            print("WARNING: heading not found:", label)
    for prefix, items in (("Table", TABLES), ("Figure", FIGURES)):
        for n, _ in items:
            key = f"{prefix} {n}:"
            for i in range(start, len(page_lines)):
                if any(l.startswith(key) for l in page_lines[i]):
                    pages[key] = i - start + 1
                    break
            else:
                print("WARNING: caption not found:", key)
    return pages, len(page_lines) - 1, pdf


if __name__ == "__main__":
    tmp = HERE / "_render" / "pass1.docx"
    tmp.parent.mkdir(exist_ok=True)
    build({}, tmp)
    pages, n1, _ = find_pages(tmp)
    build(pages, OUT)
    pages2, n2, pdf = find_pages(OUT)
    changed = {k: (pages[k], pages2.get(k)) for k in pages if pages2.get(k) != pages[k]}
    if changed:
        print("page numbers moved in pass 2, rebuilding:", changed)
        build(pages2, OUT)
        pages2, n2, pdf = find_pages(OUT)
    print(f"pages in PDF: {n2}; Introduction..end numbered 1..{max(pages2.values())}+")
    print(f"wrote {OUT}")
