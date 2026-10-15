"""Builds the PBL-3 presentation from the department template (DSE3170_PPT_Templates_for_PBL-3.pptx).

Keeps the template's title slide and Outline slide (filled in), then adds one or more slides for each
Outline item in the same look: the template's textured background, its rounded header box (Bookman Old
Style 28 bold on a light blue-grey gradient) and its pastel rounded bars with Times New Roman text.
Every number comes from crop_experiments/results/system (see results/CLAIMS.md).
Run from this folder: python3 build_ppt.py
"""
import copy
import io
import re
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "DSE3170_PPT_Templates_for_PBL-3.pptx"
OUT = HERE / "Crop_Recommendation_System_PBL3.pptx"
SYS = HERE.parent / "crop_experiments" / "results" / "system" / "figures"
REP = HERE.parent / "report" / "figures"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
HEAD_FONT, BODY_FONT = "Bookman Old Style", "Times New Roman"
INK, INK2, ACCENT = "1F2A36", "3A4652", "3B5A78"
GOOD, BAD = "2E6B45", "A0442C"
HDR = ("E8F0F3", "CBDCE4")                          # template header gradient (top, bottom)
PASTEL = [("EDEBF6", "D1CCE3"), ("E9EAF4", "D2D5E8"), ("E8EEF4", "CDD9E6"), ("E7F3F1", "CDE3DF"),
          ("E6F3EC", "CBE3D6"), ("E7F1E6", "CCE0CC"), ("EDF1E6", "D6E0CC")]   # template bar colours
LAV, PERI, BLUE, TEAL, MINT, GREEN, OLIVE = PASTEL
WHITE_CARD = ("FBFCFD", "EEF3F6")
X0, W0 = 0.49, 9.0                                   # content left edge and width (inches)


# ----------------------------------------------------------------------------- low-level helpers
def rgb(h):
    return RGBColor.from_string(h)


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def gradient(shape, top, bottom):
    sppr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:pattFill", "a:blipFill"):
        for e in sppr.findall(qn(tag)):
            sppr.remove(e)
    g = etree.fromstring(
        f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{top}"/></a:gs><a:gs pos="100000"><a:srgbClr val="{bottom}"/></a:gs>'
        f'</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>')
    sppr.find(qn("a:prstGeom")).addnext(g)


def shadow(shape):
    e = etree.fromstring(
        f'<a:effectLst xmlns:a="{A_NS}"><a:outerShdw blurRad="40000" dist="20000" dir="5400000" '
        f'rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="38000"/></a:srgbClr></a:outerShdw></a:effectLst>')
    shape._element.spPr.append(e)


TOKEN = re.compile(r"(\*\*.+?\*\*)")


def add_runs(p, text, size, font=BODY_FONT, color=INK, bold=False, italic=False):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        b = bold
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        r = p.add_run()
        r.text = tok
        f = r.font
        f.name, f.size, f.bold, f.italic = font, Pt(size), b, italic
        f.color.rgb = rgb(color)


def bullet(p, level_indent=0.22):
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", str(int(Inches(level_indent)))); ppr.set("indent", str(-int(Inches(level_indent))))
    bf = etree.SubElement(ppr, qn("a:buFont")); bf.set("typeface", "Arial")
    bc = etree.SubElement(ppr, qn("a:buChar")); bc.set("char", "•")


def fill_text(shape, paras, size=16, font=BODY_FONT, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
              bullets=False, space_after=4, margins=(0.12, 0.06), bold=False, italic=False, line_spacing=None):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margins[0])
    tf.margin_top = tf.margin_bottom = Inches(margins[1])
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = {}
        if isinstance(item, tuple):
            item, opts = item
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        if opts.get("bullet", bullets):
            bullet(p)
        add_runs(p, item, opts.get("size", size), opts.get("font", font), opts.get("color", color),
                 opts.get("bold", bold), opts.get("italic", italic))
    return tf


def box(slide, x, y, w, h, colors=WHITE_CARD, line=None, line_w=1.0, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        shade=True, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(s)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if colors is None:
        s.fill.background()
    elif isinstance(colors, str):
        s.fill.solid(); s.fill.fore_color.rgb = rgb(colors)
    else:
        gradient(s, *colors)
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if shade:
        shadow(s)
    if name:
        s.name = name
    return s


def text(slide, x, y, w, h, paras, **kw):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    kw.setdefault("anchor", MSO_ANCHOR.TOP)
    kw.setdefault("margins", (0.05, 0.03))
    fill_text(t, paras, **kw)
    return t


def bar(slide, x, y, w, h, paras, colors=BLUE, size=16, **kw):
    s = box(slide, x, y, w, h, colors)
    fill_text(s, paras, size=size, **kw)
    return s


def circle_num(slide, x, y, d, n, color=ACCENT, size=16):
    c = box(slide, x, y, d, d, colors=color, shape=MSO_SHAPE.OVAL, shade=False)
    fill_text(c, str(n), size=size, font=HEAD_FONT, color="FFFFFF", align=PP_ALIGN.CENTER, bold=True,
              margins=(0, 0))
    return c


def arrow(slide, x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, color="7A93A8"):
    return box(slide, x, y, w, h, colors=color, shape=kind, shade=False)


def picture(slide, path, x, y, maxw, maxh, center=True):
    iw, ih = Image.open(path).size
    scale = min(maxw / iw, maxh / ih)
    w, h = iw * scale, ih * scale
    px = x + (maxw - w) / 2 if center else x
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = rgb("B8C4CE"); pic.line.width = Pt(0.75)
    return pic, (px, y, w, h)


def table(slide, x, y, widths, rows, size=13, header_fill=HDR[1], highlight=(), row_h=0.3, align=None,
          bold_first_col=False):
    nr, nc = len(rows), len(widths)
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * nr))
    tbl = gf.table
    tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            elif r in highlight:
                cell.fill.fore_color.rgb = rgb("DCEBDD")
            else:
                cell.fill.fore_color.rgb = rgb("FFFFFF" if r % 2 else "F1F5F8")
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            a = PP_ALIGN.LEFT
            if align and c < len(align) and align[c] == "c":
                a = PP_ALIGN.CENTER
            p.alignment = PP_ALIGN.CENTER if r == 0 else a
            add_runs(p, rows[r][c], size, BODY_FONT, INK, bold=(r == 0 or (bold_first_col and c == 0)))
    return gf


# ----------------------------------------------------------------------------- slide scaffolding
class Deck:
    def __init__(self):
        self.prs = Presentation(str(TEMPLATE))
        self.title_slide, self.outline_slide = self.prs.slides[0], self.prs.slides[1]
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "Blank")
        bg = self.title_slide._element.find(qn("p:cSld")).find(qn("p:bg"))
        self.bg_xml = copy.deepcopy(bg)
        rid = bg.find(".//" + qn("a:blip")).get(qn("r:embed"))
        self.bg_part = self.title_slide.part.related_part(rid)
        self.sldnum = copy.deepcopy(next(sp for sp in self.title_slide.shapes._spTree.iter(qn("p:sp"))
                                         if sp.find(".//" + qn("p:ph")) is not None
                                         and sp.find(".//" + qn("p:ph")).get("type") == "sldNum"))

    def new(self, title, notes):
        s = self.prs.slides.add_slide(self.blank)
        csld = s._element.find(qn("p:cSld"))
        bg = copy.deepcopy(self.bg_xml)
        rid = s.part.relate_to(self.bg_part, RT.IMAGE)
        bg.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
        csld.insert(0, bg)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        if title:
            h = box(s, 0.488, 0.134, 9.0, 0.95, colors=HDR, line="FFFFFF", line_w=2.0, radius=0.16, shade=False,
                    name="Title")
            fill_text(h, title, size=28, font=HEAD_FONT, color="000000", align=PP_ALIGN.CENTER, bold=True)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def finish_numbers(self):
        for s in list(self.prs.slides)[2:]:
            sp = copy.deepcopy(self.sldnum)
            tree = s.shapes._spTree
            ids = [int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))]
            sp.find(".//" + qn("p:cNvPr")).set("id", str(max(ids) + 1))
            tree.append(sp)


# ----------------------------------------------------------------------------- template slides
def fill_title(d):
    s = d.title_slide
    for shp in s.shapes:
        if shp.has_text_frame and "Project Based Learning Title" in shp.text_frame.text:
            p = shp.text_frame.paragraphs[0]
            p.runs[0].text = "Crop Recommendation System"
            sub = copy.deepcopy(p._p)
            p._p.addnext(sub)
            r = sub.find(qn("a:r"))
            r.find(qn("a:t")).text = "Using Machine Learning, Explainable AI and Fertilizer Recommendation"
            rpr = r.find(qn("a:rPr")); rpr.set("sz", "1600"); rpr.set("b", "0")
        if shp.has_table:
            tbl = shp.table
            c0 = tbl.cell(0, 0).text_frame
            p = c0.paragraphs[1]
            p.alignment = PP_ALIGN.LEFT
            add_runs(p, "[Student 1 name] ([Reg. No.])", 14)
            for line in ("[Student 2 name] ([Reg. No.])", "[Student 3 name] ([Reg. No.])"):
                q = c0.add_paragraph(); q.alignment = PP_ALIGN.LEFT
                add_runs(q, line, 14)
            c2 = tbl.cell(0, 2).text_frame
            for line in ("[Guide name]", "[Designation]", "Dept. of Data Science and Engineering"):
                q = c2.add_paragraph(); q.alignment = PP_ALIGN.LEFT
                add_runs(q, line, 14)
    s.notes_slide.notes_text_frame.text = (
        "Introduce the team and the guide. One line on the project: a system that tells a farmer which crop to "
        "grow from a soil test and the weather, explains why, and suggests a fertilizer. Replace the bracketed "
        "names before presenting.")
    d.outline_slide.notes_slide.notes_text_frame.text = (
        "The talk follows the department outline: abstract, introduction, motivation, objective, literature "
        "survey, research gap, problem statement, proposed methodology, results and outcome, and references.")


# ----------------------------------------------------------------------------- content slides
def slides(d):
    # Abstract ------------------------------------------------------------------------------
    s = d.new("Abstract", "Summarise the whole project in one minute: the problem, what we built, the main "
              "numbers. Crop model: Random Forest and Naive Bayes both 99.55% on 440 unseen samples. Fertilizer "
              "model: XGBoost 97.69%. Explanations: humidity and rainfall matter most.")
    bar(s, X0, 1.3, W0, 3.75, [
        ("Farmers often choose crops by habit. Soil tests and weather data are available, but turning them into "
         "a decision is hard.", {"bullet": True}),
        ("We built a system that recommends crops from 7 inputs (N, P, K, pH, temperature, humidity, rainfall), "
         "explains each recommendation and suggests a fertilizer.", {"bullet": True}),
        ("Ten machine learning algorithms were trained and tested on 2,200 samples of 22 crops. Random Forest and "
         "Naive Bayes were best: **99.55%** (438 of 440 test samples).", {"bullet": True}),
        ("SHAP and LIME show that humidity and rainfall matter most, followed by N, P and K.", {"bullet": True}),
        ("For fertilizer, XGBoost reached **97.69%** accuracy; a nutrient-gap rule tells the farmer which "
         "nutrient to add or skip.", {"bullet": True})], colors=LAV, size=17, space_after=8)
    for i, (num, lab, col) in enumerate([("99.55%", "crop accuracy (Random Forest)", BLUE),
                                          ("97.69%", "fertilizer accuracy (XGBoost)", TEAL),
                                          ("22 crops", "10 algorithms compared", MINT)]):
        b = box(s, X0 + i * 3.1, 5.3, 2.8, 1.5, col)
        fill_text(b, [(num, {"size": 30, "font": HEAD_FONT, "bold": True, "color": ACCENT}),
                      (lab, {"size": 14})], align=PP_ALIGN.CENTER, space_after=2)

    # Introduction --------------------------------------------------------------------------
    s = d.new("Introduction", "Choosing a crop is the first decision of every season. A soil test gives N, P, K "
              "and pH, and the weather gives temperature, humidity and rainfall. Machine learning can learn which "
              "crop suits such a field. The diagram is the whole system: the farmer's input goes to the crop "
              "model, which gives the top 3 crops; then the explanation and the fertilizer advice.")
    bar(s, X0, 1.3, W0, 0.62, "Choosing the right crop for the soil and the season decides a farmer's income "
        "for the whole year.", colors=BLUE)
    bar(s, X0, 2.02, W0, 0.62, "Machine learning can learn from soil (N, P, K, pH) and weather (temperature, "
        "humidity, rainfall) which crop suits a field [1].", colors=TEAL)
    picture(s, REP / "fig1_1_architecture.png", X0, 2.82, W0, 3.75)
    text(s, X0, 6.6, W0, 0.3, "Proposed system: recommend the crop, explain why, advise on fertilizer",
         size=13, italic=True, align=PP_ALIGN.CENTER, color=INK2)

    # Motivation ----------------------------------------------------------------------------
    s = d.new("Motivation", "Three reasons for the project. Farmers often decide by habit. The data needed for a "
              "better decision exists but is hard to use. And a recommendation with no reason is hard to trust; "
              "the next question is always which fertilizer to use.")
    cards = [("Crop choice by habit", "Farmers often sow what neighbours grow or what paid well last season, "
              "even when soil and rainfall have changed.", LAV),
             ("The data exists", "Soil Health Card tests give N, P, K and pH, and weather data is public. Turning "
              "these numbers into a decision is the hard part.", BLUE),
             ("Advice needs reasons", "A crop name with no explanation is hard to trust, and the next question is "
              "always: which fertilizer?", MINT)]
    for i, (t, body, col) in enumerate(cards):
        x = X0 + i * 3.1
        box(s, x, 1.35, 2.8, 3.95, col)
        circle_num(s, x + 0.2, 1.55, 0.55, i + 1)
        text(s, x + 0.2, 2.25, 2.4, 0.9, t, size=18, font=HEAD_FONT, bold=True)
        text(s, x + 0.2, 3.15, 2.4, 2.0, body, size=15)
    bar(s, X0, 5.55, W0, 1.1, "**Our motivation:** give the farmer advice they can check and act on: the crop, "
        "the reason behind it, and the fertilizer.", colors=OLIVE, size=17)

    # Objective -----------------------------------------------------------------------------
    s = d.new("Objective", "Six objectives. Read them in order; each one maps to a part of the results later.")
    objs = ["Recommend the most suitable crops (top 3, with probabilities) from soil N, P, K, pH and the weather.",
            "Compare 10 machine learning algorithms on the same train/test split and choose the best with clear "
            "reasons.",
            "Check whether the high accuracy of the chosen model is a sign of overfitting.",
            "Explain every recommendation with SHAP and LIME.",
            "Recommend a fertilizer: predict a product and compare the soil with the crop's needs.",
            "Write up the work as a research paper in IEEE conference format."]
    for i, o in enumerate(objs):
        y = 1.3 + i * 0.92
        b = bar(s, X0, y, W0, 0.78, o, colors=PASTEL[i], size=16, margins=(0.8, 0.05))
        circle_num(s, X0 + 0.15, y + 0.14, 0.5, i + 1)

    # Literature survey ---------------------------------------------------------------------
    s = d.new("Literature Survey", "Earlier work falls into three groups: reviews of machine learning in "
              "agriculture, crop recommenders, and methods for explaining models. Harvestify is the open-source "
              "project our crop dataset comes from.")
    rows = [["Ref.", "Study", "Work", "Key idea / relevance"],
            ["[1]", "Liakos et al. (2018)", "Review: ML in agriculture", "Groups work into crop, livestock, water "
             "and soil management"],
            ["[2]", "van Klompenburg et al. (2020)", "Review: crop yield prediction", "Temperature, rainfall and "
             "soil type among the most used features"],
            ["[3]", "Kumar et al. (2015)", "Crop selection method", "Chooses a crop sequence to improve the yield "
             "rate"],
            ["[4]", "Pudumalar et al. (2017)", "Ensemble crop recommender", "Combines several classifiers on soil "
             "data"],
            ["[5]", "Doshi et al. (2018)", "AgroConsultant", "Compares classifiers using soil and climate inputs"],
            ["[6]", "Harvestify (GitHub)", "Crop and fertilizer web app", "Source of our crop dataset and crop "
             "nutrient table"],
            ["[7]", "Lundberg and Lee (2017)", "SHAP", "Splits a prediction into feature contributions"],
            ["[8]", "Ribeiro et al. (2016)", "LIME", "Explains one prediction with a local linear model"],
            ["[9]", "Ryo (2022)", "Explainable AI in agriculture", "Explainable ML for agricultural data analysis"]]
    table(s, X0, 1.3, [0.55, 2.35, 2.3, 3.8], rows, size=13, align="c", row_h=0.5)

    # Research gap --------------------------------------------------------------------------
    s = d.new("Research Gap", "For each gap in earlier work, the right column says what our project does about "
              "it.")
    text(s, X0, 1.3, 4.0, 0.4, "What earlier work does", size=17, font=HEAD_FONT, bold=True, align=PP_ALIGN.CENTER)
    text(s, X0 + 5.0, 1.3, 4.0, 0.4, "What our project does", size=17, font=HEAD_FONT, bold=True,
         align=PP_ALIGN.CENTER)
    gaps = [("Models compared on different splits, usually by accuracy only",
             "10 algorithms on the same split: accuracy, precision, recall, F1, cross-validation, time and an "
             "overfitting check"),
            ("A crop name is given with no reason",
             "SHAP and LIME explain every recommendation: overall, per crop and per farmer"),
            ("Fertilizer advice built as a separate project",
             "A fertilizer module linked to the recommended crop: ML model plus nutrient-gap rule")]
    for i, (a, b) in enumerate(gaps):
        y = 1.85 + i * 1.6
        bar(s, X0, y, 4.0, 1.35, a, colors=LAV, size=16)
        arrow(s, X0 + 4.2, y + 0.45, 0.6, 0.45)
        bar(s, X0 + 5.0, y, 4.0, 1.35, b, colors=GREEN, size=16)

    # Problem statement ---------------------------------------------------------------------
    s = d.new("Problem Statement", "Read the statement, then walk through the strip: seven inputs go in, three "
              "kinds of output come out.")
    bar(s, X0, 1.35, W0, 2.2, "Given the N, P, K and pH of a soil sample and the temperature, humidity and "
        "rainfall of the area, recommend the crops that best suit these conditions, explain which inputs led to "
        "the recommendation, and suggest how to correct the soil with fertilizer.", colors=BLUE, size=20,
        margins=(0.3, 0.1))
    y, h = 4.05, 2.6
    b = box(s, X0, y, 2.6, h, LAV)
    fill_text(b, [("Inputs", {"size": 18, "font": HEAD_FONT, "bold": True, "align": PP_ALIGN.CENTER}),
                  ("Soil test: N, P, K, pH", {"bullet": True}), ("Weather: temperature, humidity, rainfall",
                                                                 {"bullet": True})], size=15)
    arrow(s, X0 + 2.7, y + 1.05, 0.45, 0.5)
    b = box(s, X0 + 3.25, y, 2.5, h, HDR, line="FFFFFF", line_w=1.5)
    fill_text(b, [("Crop Recommendation System", {"size": 16, "font": HEAD_FONT, "bold": True}),
                  ("ML model, XAI and fertilizer module", {"size": 14})], align=PP_ALIGN.CENTER)
    arrow(s, X0 + 5.85, y + 1.05, 0.45, 0.5)
    b = box(s, X0 + 6.4, y, 2.6, h, GREEN)
    fill_text(b, [("Outputs", {"size": 18, "font": HEAD_FONT, "bold": True, "align": PP_ALIGN.CENTER}),
                  ("Top 3 crops with probabilities", {"bullet": True}), ("Reasons (SHAP, LIME)", {"bullet": True}),
                  ("Fertilizer advice", {"bullet": True})], size=15)

    # Proposed methodology: overview ---------------------------------------------------------
    s = d.new("Proposed Methodology", "Six steps, in the order of the numbers: data, pre-processing, training, "
              "testing, explanation, fertilizer advice. The next slides go through each step.")
    steps = [("Data collection", "3 public datasets: crops (2,200 rows), fertilizer (99 rows), crop nutrient "
              "needs (22 crops)"),
             ("Pre-processing", "Label encoding, stratified 80/20 split (1,760 / 440), scaling for KNN, SVM, LR"),
             ("Training", "10 algorithms, default settings, 5-fold cross-validation on the training data"),
             ("Testing", "Accuracy, precision, recall, F1 on 440 unseen samples; overfitting check"),
             ("Explanation", "SHAP (TreeSHAP) and LIME on the Random Forest"),
             ("Fertilizer advice", "XGBoost classifier plus nutrient-gap rule for the chosen crop")]
    pos = [(0, 0), (1, 0), (2, 0), (2, 1), (1, 1), (0, 1)]
    for i, ((t, body), (col, row)) in enumerate(zip(steps, pos)):
        x, y = X0 + col * 3.15, 1.35 + row * 2.85
        box(s, x, y, 2.7, 2.35, PASTEL[i])
        circle_num(s, x + 0.15, y + 0.15, 0.5, i + 1)
        text(s, x + 0.75, y + 0.2, 1.85, 0.45, t, size=16, font=HEAD_FONT, bold=True)
        text(s, x + 0.15, y + 0.8, 2.4, 1.5, body, size=14)
    arrow(s, X0 + 2.75, 2.3, 0.35, 0.4); arrow(s, X0 + 5.9, 2.3, 0.35, 0.4)
    arrow(s, X0 + 7.8, 3.75, 0.4, 0.42, kind=MSO_SHAPE.DOWN_ARROW)
    arrow(s, X0 + 5.9, 5.15, 0.35, 0.4, kind=MSO_SHAPE.LEFT_ARROW)
    arrow(s, X0 + 2.75, 5.15, 0.35, 0.4, kind=MSO_SHAPE.LEFT_ARROW)

    # Methodology: datasets -------------------------------------------------------------------
    s = d.new("Methodology: Datasets", "Three public datasets. The crop dataset trains the crop model, the "
              "fertilizer dataset trains the fertilizer model, and the nutrient table drives the nutrient-gap "
              "rule. Links: paper/DATASETS.md.")
    rows = [["Dataset", "Source", "Rows", "Classes", "Inputs"],
            ["Crop Recommendation", "Kaggle [10]", "2,200", "22 crops, 100 each",
             "N, P, K, temperature, humidity, pH, rainfall"],
            ["Fertilizer Prediction", "Kaggle [11]", "99", "7 fertilizers",
             "Temperature, humidity, moisture, soil type, crop type, N, P, K"],
            ["Crop nutrient needs", "Harvestify [6]", "22", "one row per crop", "Required N, P, K and pH"]]
    table(s, X0, 1.35, [2.1, 1.45, 0.75, 1.7, 3.0], rows, size=14, align="lcccl", row_h=0.55, bold_first_col=True)
    bar(s, X0, 4.25, W0, 0.62, "Balanced: exactly 100 samples per crop, no missing values, no duplicates",
        colors=BLUE)
    bar(s, X0, 5.0, W0, 0.62, "Fertilizers: Urea 22, DAP 18, 28-28 17, 14-35-14 14, 20-20 14, 17-17-17 7, "
        "10-26-26 7 rows", colors=TEAL)
    bar(s, X0, 5.75, W0, 0.62, "All three are public, so anyone can repeat our experiments with our code",
        colors=MINT)

    # Methodology: data analysis -------------------------------------------------------------
    s = d.new("Methodology: Data Analysis", "Left: correlation between the seven inputs; only P and K move "
              "together. Right: average N, P, K per crop; apple and grapes need far more P and K, cotton the "
              "most N. Then the split: 80% for training, 20% for testing, 20 test samples per crop.")
    picture(s, SYS / "figS1_data_analysis.png", X0, 1.3, W0, 3.85)
    bar(s, X0, 5.3, W0, 0.46, "P and K are strongly related (r = 0.74); all other pairs are weak (−0.23 to "
        "0.21)", colors=LAV, size=14)
    bar(s, X0, 5.84, W0, 0.46, "Apple and grapes need the most P and K; cotton needs the most N (about 118)",
        colors=BLUE, size=14)
    bar(s, X0, 6.38, W0, 0.46, "Stratified 80/20 split: 1,760 training and 440 test samples (20 per crop); "
        "scaling for KNN, SVM, LR", colors=MINT, size=14)

    # Methodology: algorithms ----------------------------------------------------------------
    s = d.new("Methodology: Algorithms", "Ten algorithms from the main families. KNN works by proximity: it "
              "votes among the five most similar training samples. Random Forest, XGBoost, LightGBM and AdaBoost "
              "combine many trees. AdaBoost is tested in two forms because the setting matters a lot.")
    algos = [("Logistic Regression", "Straight boundary for each crop"),
             ("Naive Bayes", "Each crop's average and spread per input"),
             ("KNN (k = 5)", "Vote of the 5 closest samples"),
             ("SVM (RBF)", "Widest gap between crops; can curve"),
             ("Decision Tree", "Yes/no questions on one input at a time"),
             ("Random Forest (chosen)", "100 trees on random samples; majority vote [12]"),
             ("XGBoost", "Adds trees that fix earlier errors [13]"),
             ("LightGBM", "Fast gradient boosting"),
             ("AdaBoost (default)", "50 one-split stumps"),
             ("AdaBoost (depth-3)", "200 small trees")]
    for i, (t, body) in enumerate(algos):
        col, row = i % 5, i // 5
        x, y = X0 + col * 1.83, 1.35 + row * 2.4
        chosen = t.startswith("Random Forest")
        box(s, x, y, 1.68, 2.2, GREEN if chosen else PASTEL[(i % 5) + (1 if row else 0)],
            line=ACCENT if chosen else None, line_w=2.25)
        text(s, x + 0.08, y + 0.12, 1.52, 0.8, t, size=14, font=HEAD_FONT, bold=True)
        text(s, x + 0.08, y + 0.95, 1.52, 1.2, body, size=14)
    bar(s, X0, 6.2, W0, 0.6, "Same 1,760 training and 440 test samples for all ten; default settings, no tuning",
        colors=OLIVE, size=15)

    # Methodology: XAI and fertilizer --------------------------------------------------------
    s = d.new("Methodology: XAI and Fertilizer", "Left: how the explanations are made. Right: the two parts of "
              "the fertilizer module, the ML model and the nutrient-gap rule.")
    b = box(s, X0, 1.35, 4.35, 5.35, BLUE)
    fill_text(b, [("Explainable AI", {"size": 20, "font": HEAD_FONT, "bold": True}),
                  ("SHAP [7]: splits each prediction into a contribution from every input; exact for tree models "
                   "(TreeSHAP [14])", {"bullet": True}),
                  ("Global importance, importance per crop, and the explanation for one farmer", {"bullet": True}),
                  ("LIME [8]: fits a simple local model around one farmer's input (2,000 perturbed samples)",
                   {"bullet": True}),
                  ("Agreement between SHAP and LIME measured on 100 test samples", {"bullet": True})],
              size=17, anchor=MSO_ANCHOR.TOP, space_after=14, margins=(0.22, 0.18))
    b = box(s, X0 + 4.65, 1.35, 4.35, 5.35, MINT)
    fill_text(b, [("Fertilizer recommendation", {"size": 20, "font": HEAD_FONT, "bold": True}),
                  ("ML model: 8 algorithms on 99 rows; 5-fold cross-validation repeated 10 times (50 tests)",
                   {"bullet": True}),
                  ("Nutrient-gap rule: gap = crop needs − soil has, for N, P and K", {"bullet": True}),
                  ("The biggest gap decides: soil short of it → add urea (N), DAP or SSP (P), MOP (K); soil "
                   "has extra → skip that nutrient", {"bullet": True})],
              size=17, anchor=MSO_ANCHOR.TOP, space_after=14, margins=(0.22, 0.18))

    # Results: table ----------------------------------------------------------------------------
    s = d.new("Result: Algorithm Comparison", "Test accuracy is on the 440 samples the models never saw. Random "
              "Forest and Naive Bayes tie at 99.55%, 2 mistakes each. The top five are within three test samples. "
              "AdaBoost with default settings is the one failure, explained two slides later.")
    rows = [["Algorithm", "Test accuracy", "F1-score", "Errors", "5-fold CV accuracy", "Training time"],
            ["Random Forest (chosen)", "99.55%", "99.55%", "2", "99.38 ± 0.61%", "0.24 s"],
            ["Naive Bayes (tied best)", "99.55%", "99.54%", "2", "99.49 ± 0.42%", "<0.01 s"],
            ["AdaBoost (depth-3 trees)", "99.32%", "99.30%", "3", "98.75 ± 0.29%", "0.92 s"],
            ["XGBoost", "99.09%", "99.08%", "4", "99.26 ± 0.14%", "1.02 s"],
            ["LightGBM", "98.86%", "98.86%", "5", "99.03 ± 0.39%", "1.44 s"],
            ["SVM (RBF)", "98.41%", "98.40%", "7", "97.90 ± 1.03%", "0.16 s"],
            ["Decision Tree", "97.95%", "97.94%", "9", "98.52 ± 0.68%", "0.01 s"],
            ["KNN (k = 5)", "97.95%", "97.93%", "9", "96.53 ± 1.21%", "<0.01 s"],
            ["Logistic Regression", "97.27%", "97.25%", "12", "96.82 ± 0.66%", "0.04 s"],
            ["AdaBoost (default stumps)", "17.50%", "8.56%", "363", "31.93 ± 5.85%", "0.14 s"]]
    table(s, X0, 1.3, [2.8, 1.25, 1.0, 0.8, 1.85, 1.3], rows, size=13, align="lccccc", highlight=(1, 2), row_h=0.37)
    bar(s, X0, 5.75, W0, 0.95, "**Random Forest and Naive Bayes tie at 99.55%** (2 errors in 440). The top five "
        "algorithms are within 3 test samples of each other.", colors=GREEN, size=16)

    # Results: chart + why RF -------------------------------------------------------------------
    s = d.new("Why We Chose Random Forest", "Accuracy alone cannot separate the top models, so we chose on four "
              "other points. Naive Bayes is an equally accurate backup, but it cannot be explained exactly.")
    picture(s, SYS / "figS2_model_comparison.png", X0, 1.3, W0, 3.6)
    reasons = ["**Joint best:** 99.55% on 440 test samples", "**Stable:** 99.38 ± 0.61% in cross-validation",
               "**Top-3 list:** tree votes give every crop a probability",
               "**Explainable:** TreeSHAP is exact for trees; trains in 0.24 s"]
    for i, r in enumerate(reasons):
        bar(s, X0 + (i % 2) * 4.58, 5.05 + (i // 2) * 0.95, 4.42, 0.82, r, colors=PASTEL[i + 2], size=15)

    # Results: AdaBoost ------------------------------------------------------------------------
    s = d.new("Why AdaBoost (Default) Failed", "The default AdaBoost uses stumps, trees with one "
              "split. One split can only divide 22 crops into two groups, so boosting adds very little. With "
              "depth-3 trees the same algorithm reaches 99.32%. A low AdaBoost score here is a setting problem, "
              "not a boosting problem.")
    for i, (num, lab, col, c) in enumerate([("17.50%", "Default: 50 one-split stumps", LAV, BAD),
                                             ("99.32%", "200 trees of depth 3", GREEN, GOOD)]):
        b = box(s, X0 + i * 4.65, 1.4, 4.35, 2.4, col)
        fill_text(b, [(num, {"size": 44, "font": HEAD_FONT, "bold": True, "color": c}), (lab, {"size": 17})],
                  align=PP_ALIGN.CENTER)
    bar(s, X0, 4.1, W0, 2.6, [
        ("A stump makes one split, so it can only divide 22 crops into 2 groups.", {"bullet": True}),
        ("Each boosting round adds very little, and most crops are never separated.", {"bullet": True}),
        ("A depth-3 tree separates up to 8 groups, and the same algorithm then works.", {"bullet": True}),
        ("**Lesson:** the weak learner was too weak, not boosting itself.", {"bullet": True})],
        colors=BLUE, size=17, space_after=8)

    # Results: confusion matrix ----------------------------------------------------------------
    s = d.new("Result: Testing Random Forest", "Rows are the true crop, columns the predicted crop. Only two "
              "mistakes, both between crops that really are similar in the data.")
    picture(s, SYS / "figS3_confusion_best.png", X0, 1.3, 4.75, 5.45, center=False)
    b = box(s, 5.45, 1.3, 4.04, 1.45, HDR, line="FFFFFF", line_w=1.5)
    fill_text(b, [("438 / 440", {"size": 32, "font": HEAD_FONT, "bold": True, "color": ACCENT}),
                  ("test samples correct; 20 of 22 crops perfect", {"size": 14})], align=PP_ALIGN.CENTER)
    bar(s, 5.45, 2.95, 4.04, 1.75, [("Mistake 1: rice → jute", {"bold": True, "font": HEAD_FONT}),
                                     "Same N, P, K (about 80-47-40) and humidity near 80%; they differ mainly in "
                                     "rainfall (236 vs 175 mm)."], colors=LAV, size=14)
    bar(s, 5.45, 4.9, 4.04, 1.5, [("Mistake 2: blackgram → maize", {"bold": True, "font": HEAD_FONT}),
                                   "Same average humidity (65%) and similar potassium."], colors=MINT, size=14)

    # Results: overfitting ---------------------------------------------------------------------
    s = d.new("Is 99.5% Overfitting? No", "Overfitting means high training accuracy and much lower test "
              "accuracy. Here the gap is under one point for most models, and the depth curve shows training and "
              "test rising together. With shuffled labels the model drops to chance.")
    picture(s, REP / "fig2_5_overfitting.png", X0, 1.3, W0, 3.65)
    bar(s, X0, 5.1, W0, 0.52, "Training vs test gap: 0.45 points for Random Forest; largest 2.05 (single decision "
        "tree)", colors=LAV, size=14)
    bar(s, X0, 5.72, W0, 0.52, "Shuffled labels: Random Forest drops to 4.61%, the chance level for 22 crops "
        "(4.55%)", colors=BLUE, size=14)
    bar(s, X0, 6.34, W0, 0.52, "Accuracy is high because the crops are well separated in this dataset",
        colors=MINT, size=14)

    # Results: farmers -------------------------------------------------------------------------
    s = d.new("Result: Six Farmers' Top-3 Crops", "Six test samples treated as six farmers. All six first "
              "choices are right. For F2 the model is less sure, so the second and third choices matter.")
    rows = [["Farmer", "Soil N-P-K", "Temp., humidity, rain", "1st choice", "2nd choice", "3rd choice"],
            ["F1", "18-17-31", "31.7 °C, 45.2%, 94 mm", "mango (0.98)", "orange (0.01)", "pigeon peas (0.01)"],
            ["F2", "22-44-24", "24.3 °C, 56.3%, 59 mm", "moth beans (0.72)", "pigeon peas (0.10)",
             "maize (0.08)"],
            ["F3", "3-67-24", "17.0 °C, 19.9%, 103 mm", "kidney beans (0.99)", "pigeon peas (0.01)",
             "blackgram (0.00)"],
            ["F4", "34-75-24", "23.5 °C, 51.3%, 192 mm", "pigeon peas (0.98)", "moth beans (0.02)",
             "blackgram (0.00)"],
            ["F5", "76-54-45", "24.3 °C, 77.6%, 185 mm", "jute (0.91)", "rice (0.08)", "papaya (0.01)"],
            ["F6", "117-76-47", "25.6 °C, 77.4%, 93 mm", "banana (0.99)", "cotton (0.01)",
             "blackgram (0.00)"]]
    table(s, X0, 1.3, [0.8, 1.1, 1.95, 1.75, 1.7, 1.7], rows, size=13, align="cc", row_h=0.55)
    bar(s, X0, 5.55, W0, 1.1, "All six first choices match the crop actually grown. When the model is less sure "
        "(F2: 0.72), the 2nd and 3rd choices give real alternatives.", colors=GREEN, size=16)

    # Results: SHAP global ---------------------------------------------------------------------
    s = d.new("Result: Which Inputs Matter Most", "SHAP over all test samples: humidity and rainfall matter most, "
              "then K, N and P. The right panel shows the deciding input changes by crop, for example nitrogen "
              "for cotton and phosphorus for apple and grapes.")
    picture(s, SYS / "figS4_shap_global_per_crop.png", X0, 1.3, W0, 4.1)
    bar(s, X0, 5.55, W0, 0.6, "Overall: humidity 22%, rainfall 18%, K 17%, N 17%, P 16%, temperature 7%, pH 3%",
        colors=LAV, size=15)
    bar(s, X0, 6.25, W0, 0.6, "Deciding input per crop: rainfall 8 crops, humidity 7, P 3, N 2 (cotton, maize), "
        "K 2", colors=MINT, size=15)

    # Results: one farmer ----------------------------------------------------------------------
    s = d.new("Result: Why Farmer F1 Got Mango", "SHAP and LIME for one farmer. Both put humidity first. They "
              "do not agree on every detail, so the system shows only the top two or three reasons.")
    picture(s, SYS / "figS5_farmer_explanation.png", X0, 1.3, W0, 3.45)
    bar(s, X0, 4.95, W0, 0.85, "**Farmer F1 → mango (0.98):** humidity 45.2% and rainfall 93.8 mm push most "
        "towards mango; the low N, P and K also fit.", colors=BLUE, size=15)
    bar(s, X0, 5.95, W0, 0.85, "**SHAP vs LIME:** same top reason in 66% of 100 test samples and 78% overlap in "
        "the top 3, so the farmer sees only the top 2–3 reasons.", colors=TEAL, size=15)

    # Results: fertilizer algorithms -----------------------------------------------------------
    s = d.new("Result: Fertilizer Algorithms", "Eight algorithms on the fertilizer data. With only 99 rows we "
              "averaged 50 tests. XGBoost is best at 97.69%. KNN and Naive Bayes are weak.")
    rows = [["Algorithm", "Accuracy", "F1-score"],
            ["XGBoost (chosen)", "97.69 ± 3.20%", "95.02%"], ["Decision Tree", "96.28 ± 4.00%", "92.24%"],
            ["Random Forest", "95.07 ± 4.43%", "89.22%"], ["Logistic Regression", "94.35 ± 4.56%", "87.59%"],
            ["SVM (RBF)", "88.09 ± 5.19%", "77.61%"], ["AdaBoost (depth-3)", "86.19 ± 12.09%", "74.14%"],
            ["KNN (k = 5)", "71.41 ± 10.40%", "63.13%"], ["Naive Bayes", "66.74 ± 9.59%", "60.58%"]]
    table(s, X0, 1.35, [2.45, 1.75, 1.15], rows, size=14, align="lcc", highlight=(1,), row_h=0.52)
    b = box(s, 6.15, 1.35, 3.34, 1.75, HDR, line="FFFFFF", line_w=1.5)
    fill_text(b, [("97.69%", {"size": 34, "font": HEAD_FONT, "bold": True, "color": ACCENT}),
                  ("XGBoost accuracy, averaged over 50 tests", {"size": 14})], align=PP_ALIGN.CENTER)
    bar(s, 6.15, 3.3, 3.34, 1.75, "**Why 50 tests?** Only 99 rows, so one 80/20 split leaves about 20 test rows. "
        "We ran 5-fold cross-validation 10 times.", colors=BLUE, size=14)
    bar(s, 6.15, 5.25, 3.34, 1.45, "**7 fertilizers:** Urea, DAP, 28-28, 14-35-14, 20-20, 17-17-17, 10-26-26",
        colors=MINT, size=14)

    # Results: fertilizer importance -----------------------------------------------------------
    s = d.new("Result: What Decides Fertilizer", "Permutation importance: shuffle one input and see how much "
              "accuracy drops. Only phosphorus, nitrogen and potassium matter. That makes sense: a fertilizer "
              "supplies what the soil lacks.")
    picture(s, SYS / "figS6_fertilizer.png", X0, 1.3, W0, 3.35)
    bar(s, X0, 4.9, W0, 0.85, "**Soil N, P, K decide it:** importance P 0.44, N 0.31, K 0.26; temperature, "
        "humidity, moisture, soil and crop type 0.006 or less.", colors=LAV, size=15)
    bar(s, X0, 5.95, W0, 0.85, "**Only 2 mistakes in 99:** 10-26-26 read as 14-35-14, 17-17-17 as 20-20. KNN and "
        "Naive Bayes fail because they weigh all 8 inputs equally.", colors=GREEN, size=15)

    # Results: nutrient gap --------------------------------------------------------------------
    s = d.new("Result: Nutrient-Gap Advice", "The rule compares the farmer's soil with what the crop "
              "needs. For the recommended crop the gaps are small; for a second-choice crop they are bigger. F3 is "
              "short of nitrogen, so the advice is urea.")
    bar(s, X0, 1.3, W0, 0.8, "**gap = what the crop needs − what the soil has.** The biggest gap decides: "
        "soil short → add it; soil has extra → skip it.", colors=BLUE, size=15)
    rows = [["Farmer", "Soil N-P-K", "Crop", "Needs N-P-K", "Biggest gap", "Advice"],
            ["F1", "18-17-31", "mango (1st)", "20-20-30", "P low by 3", "Add phosphate (DAP or SSP)"],
            ["F1", "18-17-31", "orange (2nd)", "20-10-10", "K high by 21", "Skip potash this season"],
            ["F3", "3-67-24", "kidney beans (1st)", "20-60-20", "N low by 17", "Add nitrogen (urea)"],
            ["F4", "34-75-24", "pigeon peas (1st)", "20-60-20", "P high by 15", "Skip phosphate this season"],
            ["F6", "117-76-47", "banana (1st)", "100-75-50", "N high by 17", "Skip nitrogen this season"],
            ["F6", "117-76-47", "cotton (2nd)", "120-40-20", "P high by 36", "Skip phosphate this season"]]
    table(s, X0, 2.3, [0.8, 1.1, 1.85, 1.2, 1.4, 2.65], rows, size=13, align="ccccc", row_h=0.45)
    bar(s, X0, 5.75, W0, 0.95, "The recommended crop usually needs little correction (F1: P short by 3). A "
        "second-choice crop needs more (F6 → cotton: 36 extra P).", colors=GREEN, size=15)

    # Outcome ----------------------------------------------------------------------------------
    s = d.new("Outcome: Paper, Patent, Copyright", "The main outcome is the conference paper. A patent "
              "does not apply because software and algorithms are excluded in India. Copyright covers our code, "
              "paper and report automatically. Fill in the paper's submission status.")
    picture(s, REP / "paper_page-1.png", X0, 1.3, 3.9, 5.45, center=False)
    outs = [("Conference paper", "8-page IEEE conference paper: “Crop Recommendation System Using Machine "
             "Learning, Explainable AI and Fertilizer Recommendation”. Status: [submitted to ___ / under "
             "review]", GREEN),
            ("Patent", "Not filed: software and algorithms are excluded under Section 3(k) of the Patents Act, "
             "1970 [16].", LAV),
            ("Copyright", "Code, paper and report are protected automatically under the Copyright Act, 1957. Data "
             "sources are credited.", BLUE)]
    for i, (t, body, col) in enumerate(outs):
        bar(s, 4.6, 1.3 + i * 1.85, 4.89, 1.65, [(t, {"bold": True, "font": HEAD_FONT, "size": 16}), body],
            colors=col, size=14)

    # Conclusion -------------------------------------------------------------------------------
    s = d.new("Conclusion and Future Scope", "Four findings on the left, four next steps on the right. The most "
              "important next step is testing on real Soil Health Card data.")
    b = box(s, X0, 1.35, 4.35, 5.35, LAV)
    fill_text(b, [("Conclusion", {"size": 20, "font": HEAD_FONT, "bold": True}),
                  ("Random Forest: 99.55% on 440 unseen samples, with a top-3 crop list", {"bullet": True}),
                  ("No overfitting: training–test gap of 0.45 points", {"bullet": True}),
                  ("SHAP: humidity and rainfall matter most; the deciding input changes by crop", {"bullet": True}),
                  ("XGBoost: 97.69% fertilizer accuracy, plus nutrient-gap advice", {"bullet": True})],
              size=18, anchor=MSO_ANCHOR.TOP, space_after=16, margins=(0.22, 0.18))
    b = box(s, X0 + 4.65, 1.35, 4.35, 5.35, MINT)
    fill_text(b, [("Future scope", {"size": 20, "font": HEAD_FONT, "bold": True}),
                  ("Test on real Soil Health Card reports with district rainfall data", {"bullet": True}),
                  ("Fertilizer amounts in kg/ha from soil-test dose rules", {"bullet": True}),
                  ("Tune the best models", {"bullet": True}),
                  ("Mobile app with explanations in the farmer's own language", {"bullet": True})],
              size=18, anchor=MSO_ANCHOR.TOP, space_after=16, margins=(0.22, 0.18))

    # References -------------------------------------------------------------------------------
    refs = [
        "K. G. Liakos, P. Busato, D. Moshou, S. Pearson and D. Bochtis, “Machine learning in agriculture: A "
        "review,” Sensors, vol. 18, no. 8, p. 2674, 2018.",
        "T. van Klompenburg, A. Kassahun and C. Catal, “Crop yield prediction using machine learning: A "
        "systematic literature review,” Computers and Electronics in Agriculture, vol. 177, p. 105709, 2020.",
        "R. Kumar, M. P. Singh, P. Kumar and J. P. Singh, “Crop selection method to maximize crop yield rate "
        "using machine learning technique,” in Proc. ICSTM, 2015, pp. 138–145.",
        "S. Pudumalar, E. Ramanujam, R. H. Rajashree, C. Kavya, T. Kiruthika and J. Nisha, “Crop "
        "recommendation system for precision agriculture,” in Proc. ICoAC, 2017, pp. 32–36.",
        "Z. Doshi, S. Nadkarni, R. Agrawal and N. Shah, “AgroConsultant: Intelligent crop recommendation "
        "system using machine learning algorithms,” in Proc. ICCUBEA, 2018, pp. 1–6.",
        "Gladiator07, “Harvestify,” GitHub repository, 2023. https://github.com/Gladiator07/Harvestify",
        "S. M. Lundberg and S.-I. Lee, “A unified approach to interpreting model predictions,” in "
        "Advances in Neural Information Processing Systems, vol. 30, 2017, pp. 4765–4774.",
        "M. T. Ribeiro, S. Singh and C. Guestrin, “‘Why should I trust you?’: Explaining the "
        "predictions of any classifier,” in Proc. ACM SIGKDD, 2016, pp. 1135–1144.",
        "M. Ryo, “Explainable artificial intelligence and interpretable machine learning for agricultural data "
        "analysis,” Artificial Intelligence in Agriculture, vol. 6, pp. 257–265, 2022.",
        "A. Ingle, “Crop Recommendation Dataset,” Kaggle. "
        "https://www.kaggle.com/datasets/atharvaingle/crop-recommendation-dataset",
        "“Fertilizer Prediction,” Kaggle. https://www.kaggle.com/datasets/gdabhishek/fertilizer-prediction",
        "L. Breiman, “Random forests,” Machine Learning, vol. 45, no. 1, pp. 5–32, 2001.",
        "T. Chen and C. Guestrin, “XGBoost: A scalable tree boosting system,” in Proc. ACM SIGKDD, 2016, "
        "pp. 785–794.",
        "S. M. Lundberg et al., “From local explanations to global understanding with explainable AI for "
        "trees,” Nature Machine Intelligence, vol. 2, no. 1, pp. 56–67, 2020.",
        "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine Learning "
        "Research, vol. 12, pp. 2825–2830, 2011.",
        "Government of India, The Patents Act, 1970 (Act No. 39 of 1970)."]
    for part, (lo, hi) in enumerate([(0, 8), (8, 16)]):
        s = d.new(f"References ({part + 1}/2)", "References in IEEE style, numbered as cited on the slides.")
        b = box(s, X0, 1.3, W0, 5.5, WHITE_CARD)
        fill_text(b, [f"[{i + 1}]  {refs[i]}" for i in range(lo, hi)], size=14, anchor=MSO_ANCHOR.TOP,
                  space_after=9, margins=(0.25, 0.2))

    # Thank you --------------------------------------------------------------------------------
    s = d.new(None, "Thank the audience and invite questions. Likely questions: is 99.5% overfitting (slide on "
              "overfitting), why Random Forest, why AdaBoost failed, and whether it works on real farms (future "
              "scope).")
    b = box(s, 1.5, 2.4, 7.0, 1.6, HDR, line="FFFFFF", line_w=2.0, radius=0.2)
    fill_text(b, [("Thank You", {"size": 44, "font": HEAD_FONT, "bold": True})], align=PP_ALIGN.CENTER,
              color="000000")
    text(s, 1.5, 4.3, 7.0, 0.6, "Questions?", size=24, font=HEAD_FONT, align=PP_ALIGN.CENTER, bold=True)
    text(s, 1.5, 5.0, 7.0, 0.5, "Crop Recommendation System · DSE3170 Project Based Learning-3", size=16,
         align=PP_ALIGN.CENTER, color=INK2)


def fix_outline(path):
    """Outline item "Introduction of" -> "Introduction" (SmartArt data + cached drawing)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ("ppt/diagrams/data2.xml", "ppt/diagrams/drawing2.xml"):
                data = data.replace(b">Introduction of <", b">Introduction<")
            zout.writestr(item, data)
    path.write_bytes(buf.getvalue())


if __name__ == "__main__":
    d = Deck()
    fill_title(d)
    slides(d)
    d.finish_numbers()
    d.prs.save(str(OUT))
    fix_outline(OUT)
    print(f"wrote {OUT} ({len(d.prs.slides)} slides)")
